from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from create_project_summary_doc import style_doc, make_flow, make_architecture
root=Path(__file__).resolve().parents[1]
make_flow(); make_architecture()
doc=Document(); style_doc(doc)
doc.sections[0].page_width=Cm(21); doc.sections[0].page_height=Cm(29.7)
for line in (root/'docs/PROJECT_OVERVIEW_FOR_ADVISOR.md').read_text(encoding='utf-8').splitlines():
    if not line.strip(): continue
    if line.startswith('# '): doc.add_paragraph(line[2:],style='Title')
    elif line.startswith('## '): doc.add_paragraph(line[3:],style='Heading 1')
    elif line.startswith('!['):
        match=re.match(r'!\[(.*?)\]\((.*?)\)',line)
        p=doc.add_paragraph(); p.paragraph_format.keep_with_next=True
        picture=p.add_run().add_picture(str((root/'docs'/match[2]).resolve()),width=Cm(17.2))
        picture._inline.docPr.set('descr',match[1])
        p=doc.add_paragraph(match[1]+' • Ảnh minh họa' if 'tmp/' in match[2] else match[1]+' • Ảnh chụp bản demo')
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].font.size=Pt(9)
    else: doc.add_paragraph(line)
doc.core_properties.title='CaseFlow AI Mo ta he thong va huong nghien cuu'
doc.save(root/'docs/TOM_TAT_DE_TAI_GUI_THAY_DO_TIEN_DUNG.docx')
print('Document generated')

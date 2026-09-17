from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "TOM_TAT_DE_TAI_GUI_THAY_DO_TIEN_DUNG.docx"
ASSETS = ROOT / "tmp" / "project_summary_assets"
ASSETS.mkdir(parents=True, exist_ok=True)

FONT_PATH = r"C:\Windows\Fonts\arial.ttf"
FONT_BOLD_PATH = r"C:\Windows\Fonts\arialbd.ttf"
BLUE = "#12355B"
CYAN = "#0E7490"
TEAL = "#0F766E"
PALE = "#EAF3F7"
GRAY = "#556270"


def font(size, bold=False):
    return ImageFont.truetype(FONT_BOLD_PATH if bold else FONT_PATH, size)


def rounded(draw, box, fill, outline=None, radius=18, width=2):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def center_text(draw, box, text, fnt, fill="#FFFFFF", spacing=5):
    bbox = draw.multiline_textbbox((0, 0), text, font=fnt, spacing=spacing, align="center")
    x = box[0] + (box[2] - box[0] - (bbox[2] - bbox[0])) / 2
    y = box[1] + (box[3] - box[1] - (bbox[3] - bbox[1])) / 2
    draw.multiline_text((x, y), text, font=fnt, fill=fill, spacing=spacing, align="center")


def arrow(draw, start, end, color=BLUE):
    draw.line([start, end], fill=color, width=5)
    x, y = end
    if end[1] > start[1]:
        draw.polygon([(x, y), (x - 9, y - 15), (x + 9, y - 15)], fill=color)
    else:
        draw.polygon([(x, y), (x - 15, y - 9), (x - 15, y + 9)], fill=color)


def make_flow():
    image = Image.new("RGB", (1600, 490), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    boxes = [
        ((55, 125, 305, 335), "Người gửi\nyêu cầu", BLUE),
        ((420, 125, 670, 335), "CaseFlow AI\ntiếp nhận", CYAN),
        ((785, 125, 1035, 335), "AI phân loại\nvà gợi ý", TEAL),
        ((1150, 125, 1400, 335), "Nhân viên xử lý\nvà phản hồi", BLUE),
    ]
    for box, label, color in boxes:
        rounded(draw, box, color, radius=22)
        center_text(draw, box, label, font(31, True))
    for a, b in [(305, 420), (670, 785), (1035, 1150)]:
        arrow(draw, (a + 20, 230), (b - 20, 230))
    draw.text((55, 40), "Luồng xử lý một yêu cầu dịch vụ", font=font(34, True), fill=BLUE)
    draw.text((55, 385), "Kết quả: yêu cầu được theo dõi, có người chịu trách nhiệm và được xử lý theo thời hạn.", font=font(25), fill=GRAY)
    path = ASSETS / "flow.png"
    image.save(path)
    return path


def make_architecture():
    image = Image.new("RGB", (1600, 720), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    layers = [
        ((80, 85, 1520, 220), "Web client\nNgười dùng | Nhân viên | Quản trị viên", "#EAF3F7", BLUE),
        ((80, 295, 1520, 430), "API service\nXác thực | Phân quyền | Quy trình yêu cầu | SLA | Thông báo", "#DFF5F2", TEAL),
        ((80, 505, 720, 670), "MongoDB\nNgười dùng | Yêu cầu | Bình luận\nTài liệu | Nhật ký", "#EAF3F7", BLUE),
        ((880, 505, 1520, 670), "AI service\nTF-IDF + Logistic Regression\nTruy xuất tri thức | LLM tùy chọn", "#DFF5F2", TEAL),
    ]
    for box, label, fill, color in layers:
        rounded(draw, box, fill, outline=color, radius=18, width=3)
        center_text(draw, box, label, font(28, True), fill=color)
    arrow(draw, (800, 220), (800, 295), CYAN)
    arrow(draw, (530, 430), (530, 505), CYAN)
    arrow(draw, (1080, 430), (1080, 505), CYAN)
    draw.text((80, 25), "Kiến trúc đề xuất", font=font(34, True), fill=BLUE)
    path = ASSETS / "architecture.png"
    image.save(path)
    return path


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def style_doc(doc):
    section = doc.sections[0]
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.85)
    section.right_margin = Cm(1.85)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.15
    for name, size in [("Title", 22), ("Heading 1", 14), ("Heading 2", 11.5)]:
        s = styles[name]
        s.font.name = "Arial"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0, 0, 0)


def add_title(doc):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run("CaseFlow AI")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("Hệ thống quản lý yêu cầu dịch vụ tích hợp AI hỗ trợ phân loại và truy xuất tri thức")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(18, 53, 91)


def add_heading(doc, text):
    doc.add_paragraph(text, style="Heading 1")


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.columns[0].width = Cm(4.25)
    table.columns[1].width = Cm(11.3)
    header = table.rows[0].cells
    for cell, text in zip(header, ["Nội dung", "Đề xuất"]):
        cell.text = text
        set_cell_shading(cell, BLUE[1:])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell)
    for i, (left, right) in enumerate(rows):
        cells = table.add_row().cells
        cells[0].text = left
        cells[1].text = right
        if i % 2 == 0:
            set_cell_shading(cells[0], "F4F7F9")
            set_cell_shading(cells[1], "F4F7F9")
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph()


def main():
    flow = make_flow()
    architecture = make_architecture()
    doc = Document()
    style_doc(doc)
    add_title(doc)

    add_heading(doc, "1. Dự án giải quyết vấn đề gì")
    doc.add_paragraph(
        "Các yêu cầu hỗ trợ nội bộ như lỗi CNTT, hành chính hoặc cơ sở vật chất thường được tiếp nhận qua điện thoại, tin nhắn và email. "
        "Dữ liệu phân tán khiến tổ chức khó theo dõi tiến độ, phân công trách nhiệm, kiểm soát thời hạn, thống kê chất lượng phục vụ và tái sử dụng kinh nghiệm xử lý."
    )
    doc.add_paragraph(
        "CaseFlow AI là hệ thống web tập trung để tiếp nhận và quản lý các yêu cầu này theo một hồ sơ xuyên suốt. AI không thay thế nhân viên; AI chỉ đọc nhanh nội dung, đưa ra gợi ý có căn cứ, còn việc xác nhận, phân công và xử lý vẫn do con người quyết định."
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(flow), width=Cm(16.2))
    p = doc.add_paragraph("Hình 1. Luồng xử lý yêu cầu đề xuất.")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(9)

    add_heading(doc, "2. Luồng nghiệp vụ và vai trò")
    doc.add_paragraph("Một yêu cầu đi qua các bước: tạo yêu cầu -> AI phân tích -> nhân viên/ quản lý kiểm tra -> phân công -> trao đổi và cập nhật tiến độ -> đóng yêu cầu -> đánh giá. Mọi thay đổi trạng thái và thao tác quan trọng đều được ghi vào timeline và nhật ký.")
    add_table(doc, [
        ("Người dùng", "Tạo yêu cầu, theo dõi trạng thái, bổ sung thông tin, trao đổi với nhân viên phụ trách và đánh giá kết quả."),
        ("Nhân viên xử lý", "Nhận phân công, cập nhật tiến độ, phản hồi, sử dụng gợi ý AI và đóng yêu cầu."),
        ("Quản trị viên", "Quản lý người dùng, danh mục, quy tắc SLA, kho tri thức, báo cáo và nhật ký hoạt động."),
        ("Quản lý", "Theo dõi hàng đợi, hồ sơ có nguy cơ trễ SLA, cụm sự cố, khối lượng nhân viên và các chỉ số vận hành."),
    ])

    doc.add_page_break()
    add_heading(doc, "3. Chức năng hệ thống")
    add_table(doc, [
        ("Quản lý yêu cầu", "Tạo yêu cầu bằng tiếng Việt, đính kèm tệp, phân loại dịch vụ, mức ưu tiên, trạng thái, thời hạn và lịch sử xử lý."),
        ("Điều phối", "Giao nhân viên, chuyển đơn vị, ghi chú nội bộ, theo dõi thời gian xử lý và cảnh báo nguy cơ vi phạm SLA."),
        ("Trao đổi", "Bình luận theo từng hồ sơ; tách ghi chú nội bộ với phản hồi hiển thị cho người gửi."),
        ("Kho tri thức", "Quản lý tài liệu theo chủ đề, trạng thái phê duyệt và phiên bản; tài liệu có thể được dùng làm nguồn truy xuất AI."),
        ("Báo cáo vận hành", "Dashboard về số lượng yêu cầu, thời gian phản hồi, tỷ lệ SLA, hàng đợi, phân bố nhóm yêu cầu và cụm sự cố."),
        ("An toàn hệ thống", "Đăng nhập JWT, phân quyền theo vai trò, băm mật khẩu, rate limit, audit log, quên mật khẩu và circuit breaker cho API AI/LLM."),
    ])

    add_heading(doc, "4. AI được tích hợp như thế nào")
    doc.add_paragraph("AI được đặt trong quy trình xử lý thay vì đứng riêng như chatbot. Mỗi dự đoán được lưu cùng độ tin cậy để người phụ trách có thể chấp nhận, sửa hoặc bỏ qua gợi ý.")
    add_table(doc, [
        ("Phân loại yêu cầu", "Đầu vào là tiêu đề và mô tả; đầu ra là nhóm dịch vụ, top nhãn và độ tin cậy. Baseline hiện tại: TF-IDF word n-gram kết hợp Logistic Regression."),
        ("Trích xuất thông tin", "Nhận dạng một số dữ liệu có cấu trúc trong câu tiếng Việt như mã sinh viên, phòng, thời hạn bằng pattern có kiểm soát."),
        ("Tìm hồ sơ tương tự", "Dùng TF-IDF và cosine similarity để tìm các yêu cầu đã xử lý gần nghĩa; hỗ trợ nhân viên học từ tình huống cũ."),
        ("RAG trên kho tri thức", "Truy xuất các đoạn tài liệu liên quan và hiển thị nguồn trích dẫn. Hiện tại ưu tiên retrieval có căn cứ; LLM là tùy chọn cho bước diễn đạt câu trả lời."),
        ("Dự báo rủi ro SLA", "Ước lượng khả năng trễ dựa trên thời gian đã xử lý, thời hạn, số lần chuyển, khối lượng và mức ưu tiên; dùng để cảnh báo quản lý."),
    ])

    doc.add_page_break()
    add_heading(doc, "5. Dữ liệu và hướng thực nghiệm AI")
    doc.add_paragraph("Bản demo đã có pipeline tái lập để chứng minh kiến trúc. Dữ liệu demo hiện chỉ phục vụ kiểm thử kỹ thuật, không được xem là kết quả khoa học hay chất lượng triển khai thực tế.")
    add_table(doc, [
        ("Dữ liệu baseline hiện có", "60 câu mẫu tiếng Việt thuộc 6 nhóm dịch vụ để service khởi động mô hình phân loại; 30 câu độc lập dùng kiểm thử pipeline; dữ liệu SLA mô phỏng có seed cố định."),
        ("Dữ liệu cần cho luận văn", "Yêu cầu đã ẩn danh, nhãn đúng do chuyên viên xác nhận, đơn vị xử lý, ưu tiên, thời gian xử lý, chuyển bộ phận, trạng thái SLA và cờ sự cố/trùng lặp."),
        ("Quy mô mục tiêu", "Tối thiểu 300-500 mẫu đã ẩn danh cho mỗi nhãn; taxonomy nên chốt 7-12 nhóm với tiêu chí vào/ra rõ ràng."),
        ("Chống rò rỉ", "Tách train/validation/test theo thời gian hoặc người gửi/nguồn; không để các yêu cầu gần trùng xuất hiện đồng thời ở train và test."),
        ("Đánh giá phân loại", "Accuracy, precision, recall, macro F1, micro F1, confusion matrix, routing accuracy và phân tích các dự đoán sai."),
        ("Đánh giá RAG và SLA", "Recall@K/Precision@K, tỷ lệ trích dẫn đúng, groundedness; SLA dùng ROC-AUC, PR-AUC, calibration và recall nhóm thực sự trễ."),
    ])
    doc.add_paragraph("Hướng nâng cấp nghiên cứu: so sánh TF-IDF word với TF-IDF character, sau đó so sánh với PhoBERT hoặc sentence-transformer tiếng Việt. Với RAG, bổ sung chunking, metadata, vector index, reranking và bộ câu hỏi kiểm thử có đáp án tham chiếu.")

    add_heading(doc, "6. Kiến trúc kỹ thuật")
    doc.add_paragraph(
        "Kiến trúc client-server tách web, API nghiệp vụ, cơ sở dữ liệu và dịch vụ AI. Cách tách này giúp thay mô hình AI mà không ảnh hưởng các chức năng quản lý yêu cầu."
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(architecture), width=Cm(16.2))
    p = doc.add_paragraph("Hình 2. Kiến trúc tổng quan của hệ thống.")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(9)

    add_table(doc, [
        ("Web", "React + TypeScript + Vite; các màn hình theo vai trò: người gửi, nhân viên, quản lý và quản trị viên."),
        ("API", "Node.js + Express + TypeScript; REST API, xác thực JWT, phân quyền RBAC, validation và quy tắc nghiệp vụ."),
        ("MongoDB", "Lưu users, organizations, cases, comments, assignments, knowledge articles, audit logs, password-reset tokens và circuit-breaker states."),
        ("AI service", "Python + FastAPI + scikit-learn; API riêng cho classification, similar cases, knowledge retrieval và SLA risk."),
        ("Tích hợp LLM", "Theo chuẩn Chat Completions, có thể cấu hình nhà cung cấp. Nếu dịch vụ ngoài lỗi hoặc quá tải, circuit breaker ngắt tạm và hệ thống vẫn chạy phần lõi."),
    ])

    doc.add_page_break()
    add_heading(doc, "7. Phần đã triển khai và cách demo")
    add_table(doc, [
        ("Đã có", "Web React, API Express, MongoDB, FastAPI AI, phân quyền, quản lý yêu cầu, bình luận, SLA, kho tri thức, dashboard vận hành, audit log, quên mật khẩu và AI baseline."),
        ("Kiểm thử hiện có", "Build thành công; 44 API tests, 13 AI tests và 6 E2E tests. Các test kiểm tra những luồng quan trọng như đăng nhập, tạo yêu cầu, phân quyền và đặt lại mật khẩu."),
        ("Demo 1", "Đăng nhập sinh viên -> tạo yêu cầu tiếng Việt -> xem AI gợi ý nhóm/độ ưu tiên/hồ sơ tương tự."),
        ("Demo 2", "Đăng nhập nhân viên hoặc quản lý -> kiểm tra/sửa gợi ý AI -> phân công -> đổi trạng thái -> thêm ghi chú nội bộ -> kiểm tra timeline và audit."),
        ("Demo 3", "Mở dashboard -> xem hồ sơ nguy cơ trễ SLA/cụm sự cố; sau đó tra kho tri thức và kiểm tra nguồn trích dẫn trước khi phản hồi."),
        ("Cách chạy", "npm install; python -m pip install -r apps/ai/requirements.txt; tạo apps/api/.env; npm run seed; npm run dev. Web: 5173, API: 4000, AI docs: 8001/docs."),
    ])

    add_heading(doc, "8. Kế hoạch hoàn thiện")
    add_table(doc, [
        ("Chốt bài toán", "Chọn bối cảnh chính là dịch vụ sinh viên và IT helpdesk trong trường đại học; đóng băng taxonomy, quy tắc SLA và tiêu chí gán nhãn."),
        ("Thu thập dữ liệu", "Tìm nguồn mở phù hợp và/hoặc lấy yêu cầu thật đã ẩn danh, có sự đồng ý của đơn vị cung cấp; lập data card và phiên bản dữ liệu."),
        ("Nâng cấp AI", "Huấn luyện baseline trên dữ liệu thật, so sánh PhoBERT/sentence-transformer, chọn ngưỡng confidence và chạy shadow mode trước khi tự động hóa."),
        ("Đánh giá sản phẩm", "Thực hiện usability testing với người dùng đại diện; đo thời gian thao tác, task success rate, SUS và phản hồi định tính."),
    ])

    add_heading(doc, "9. Các điểm cần góp ý về chuyên môn")
    doc.add_paragraph(
        "1) Chọn taxonomy và ngữ cảnh dịch vụ nào đủ hẹp để có dữ liệu chất lượng nhưng vẫn thể hiện tính tổng quát?\n"
        "2) Baseline nào và mô hình nào nên dùng để so sánh nhằm tạo phần nghiên cứu có ý nghĩa?\n"
        "3) Quy mô dữ liệu, cách gán nhãn và cách chia tập nào phù hợp trong phạm vi đồ án?\n"
        "4) Tiêu chí đánh giá RAG, SLA risk và usability nào nên ưu tiên để chứng minh hiệu quả thực tế?\n"
        "5) Ranh giới hợp lý giữa hệ thống sản phẩm hoàn chỉnh và phần nghiên cứu AI cần tập trung là gì?"
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "CaseFlow AI"
    doc.core_properties.author = "Nhom sinh vien"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
